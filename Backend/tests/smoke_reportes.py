"""
Smoke test de la reportabilidad EPP (Fase 5) — Service → Repository.

Levanta el esquema desde el ORM en una base PostgreSQL desechable, siembra datos
sintéticos y ejercita los tres reportes + el dashboard. No usa TestClient por el
choque de versión de httpx del proyecto, y no necesita el túnel a RRHH: todos
los datos se crean acá.

Uso:
    # 1. Base desechable
    docker run -d --name pg_smoke -e POSTGRES_PASSWORD=test \
        -e POSTGRES_DB=db_smoke -p 55432:5432 postgres:16-alpine

    # 2. Correr desde Backend/
    PYTHONIOENCODING=utf-8 PYTHONPATH=. venv/Scripts/python.exe tests/smoke_reportes.py

    # 3. Limpiar
    docker rm -f pg_smoke

CUIDADO: hace drop_all sobre la base destino. Nunca apuntarlo a una base real.
"""
import os
import sys

os.environ["DATABASE_URL"] = os.getenv(
    "SMOKE_DATABASE_URL",
    "postgresql+psycopg2://postgres:test@localhost:55432/db_smoke",
)

from sqlalchemy import text                                       # noqa: E402
from app.db.session import Base, SessionLocal, engine             # noqa: E402
import app.models.inventario                                      # noqa: E402,F401
from app.repositories.entregas_repository import EntregasRepository  # noqa: E402
from app.repositories.epp_repository import EppRepository         # noqa: E402
from app.services.reportes_service import ReportesService         # noqa: E402

fallos = []


def check(nombre, condicion, detalle=""):
    print(f"  [{'OK  ' if condicion else 'FALLA'}] {nombre}" + (f" — {detalle}" if detalle else ""))
    if not condicion:
        fallos.append(nombre)


def seccion(t):
    print(f"\n{'=' * 70}\n{t}\n{'=' * 70}")


def escalar(db, sql, **p):
    return db.execute(text(sql), p).scalar()


# ── Esquema y datos ──────────────────────────────────────────────────────────

seccion("0. Esquema desde el modelo ORM")
Base.metadata.drop_all(engine)
Base.metadata.create_all(engine)
db = SessionLocal()
check("tablas EPP creadas", escalar(db, """
    SELECT COUNT(*) FROM information_schema.tables
    WHERE table_name IN ('entregas_epp','stock_epp','movimientos_stock','productos_epp')
""") == 4)

seccion("1. Siembra de datos sintéticos")

db.execute(text("INSERT INTO empresa (empresa_id, nombre_empresa) VALUES (1,'Empresa Uno'),(2,'Empresa Dos')"))
db.execute(text("""INSERT INTO areas (area_id, nombre_area, empresa_id)
                   VALUES (1,'Operaciones',1),(2,'Administración',1),(3,'Operaciones',2)"""))
db.execute(text("""INSERT INTO subareas (subarea_id, nombre_subarea, area_id)
                   VALUES (1,'Planta',1),(2,'Contabilidad',2)"""))
db.execute(text("INSERT INTO tallas (talla_id, nombre_talla) VALUES (1,'M'),(2,'L')"))
db.execute(text("INSERT INTO categorias_epp (categoria_id, nombre_categoria) VALUES (1,'Cabeza'),(2,'Calzado')"))
db.execute(text("""INSERT INTO productos_epp (producto_id, nombre, categoria_id, talla_aplica, activo)
                   VALUES (1,'Casco',1,FALSE,TRUE),
                          (2,'Botas de seguridad',2,TRUE,TRUE),
                          (3,'Protector auditivo',1,FALSE,TRUE)"""))
db.execute(text("""INSERT INTO personal (rut, nombre_completo, empresa_id, cargo, area_id, subarea_id, activo)
                   VALUES ('1-1','Ana Operaria',   1,'Operaria',   1,   1,   TRUE),
                          ('2-2','Beto Contador',  1,'Contador',   2,   2,   TRUE),
                          ('3-3','Carla Sin EPP',  1,'Operaria',   1,   1,   TRUE),
                          ('4-4','Dario Egresado', 2,'Operario',   3,   NULL,FALSE)"""))
db.commit()

epp = EppRepository(db)
# Casco (sin talla) y protector: stock global; botas por talla.
epp.ajustar_stock(1, None, 20, "carga inicial", None)
epp.ajustar_stock(2, 1, 10, "carga inicial", None)
epp.ajustar_stock(2, 2, 10, "carga inicial", None)
epp.ajustar_stock(3, None, 5, "carga inicial", None)
db.execute(text("UPDATE stock_epp SET stock_minimo = 5 WHERE producto_id = 1"))
db.execute(text("UPDATE stock_epp SET stock_minimo = 4 WHERE producto_id = 2"))
db.execute(text("UPDATE stock_epp SET stock_minimo = 3 WHERE producto_id = 3"))
db.commit()
check("stock sembrado", escalar(db, "SELECT COUNT(*) FROM stock_epp") == 4)
check("casco sin talla, una sola fila",
      escalar(db, "SELECT COUNT(*) FROM stock_epp WHERE producto_id=1") == 1)

entregas = EntregasRepository(db)
ana = entregas.get_trabajador("1-1")
beto = entregas.get_trabajador("2-2")
dario = entregas.get_trabajador("4-4", solo_activos=False)

# Ana: casco NUEVA + botas NUEVA
e_casco_ana = entregas.crear_entregas(ana, [
    {"producto_id": 1, "talla_id": None, "cantidad": 1, "motivo": "NUEVA"},
    {"producto_id": 2, "talla_id": 1, "cantidad": 1, "motivo": "NUEVA"},
], None)
id_casco_ana = e_casco_ana[0]["entrega_id"]

# Beto: protector NUEVA
e_beto = entregas.crear_entregas(beto, [
    {"producto_id": 3, "talla_id": None, "cantidad": 1, "motivo": "NUEVA"},
], None)

# Dario (desvinculado) conserva un casco sin devolver
entregas.crear_entregas(dario, [
    {"producto_id": 1, "talla_id": None, "cantidad": 1, "motivo": "NUEVA"},
], None)

seccion("2. Q6 — PERDIDA vinculada retira la entrega perdida de los vigentes")
vig_antes = len(entregas.get_vigentes_por_rut("1-1"))
entregas.crear_entregas(ana, [{
    "producto_id": 1, "talla_id": None, "cantidad": 1, "motivo": "PERDIDA",
    "entrega_reemplazada_id": id_casco_ana,
}], None)
vig_despues = entregas.get_vigentes_por_rut("1-1")
check("vigentes no aumentan tras la reposición",
      len(vig_despues) == vig_antes, f"antes={vig_antes} después={len(vig_despues)}")
check("el casco perdido ya no figura vigente",
      all(v["entrega_id"] != id_casco_ana for v in vig_despues))
check("PERDIDA no genera BAJA_DANO",
      escalar(db, "SELECT COUNT(*) FROM movimientos_stock WHERE tipo='BAJA_DANO'") == 0)

seccion("3. Sustitución por daño (DANO)")
botas_ana = next(v for v in vig_despues if v["producto_id"] == 2)
entregas.crear_sustitucion(ana, botas_ana["entrega_id"], 2, 1, 1, "rotas", None, None)
vig_final = entregas.get_vigentes_por_rut("1-1")
check("la entrega dañada sale de vigentes",
      all(v["entrega_id"] != botas_ana["entrega_id"] for v in vig_final))
check("BAJA_DANO registrada",
      escalar(db, "SELECT COUNT(*) FROM movimientos_stock WHERE tipo='BAJA_DANO'") == 1)

seccion("4. Zona horaria — corte de mes en el borde")
# 2026-08-01 02:00 UTC = 2026-07-31 22:00 en Chile (UTC-4). Debe contar en JULIO.
db.execute(text("""
    UPDATE entregas_epp SET fecha_entrega = TIMESTAMP '2026-08-01 02:00:00'
    WHERE entrega_id = :id
"""), {"id": e_beto[0]["entrega_id"]})
db.commit()

svc = ReportesService(db)
julio = svc.trazabilidad({"desde": "2026-07-01", "hasta": "2026-07-31"})
agosto = svc.trazabilidad({"desde": "2026-08-01", "hasta": "2026-08-31"})
check("la entrega del borde cae en julio (hora local)",
      any(r["entrega_id"] == e_beto[0]["entrega_id"] for r in julio))
check("y NO en agosto",
      all(r["entrega_id"] != e_beto[0]["entrega_id"] for r in agosto))
fecha_local = next(r for r in julio if r["entrega_id"] == e_beto[0]["entrega_id"])["fecha_entrega"]
check("fecha convertida a hora de Chile", fecha_local.day == 31 and fecha_local.hour == 22,
      str(fecha_local))

seccion("5. R1 — Trazabilidad")
todas = svc.trazabilidad({})
check("una fila por entrega",
      len(todas) == escalar(db, "SELECT COUNT(*) FROM entregas_epp"), f"{len(todas)} filas")
check("trae ficha organizacional",
      all(r["empresa"] for r in todas) and any(r["area"] for r in todas))
check("filtro por motivo", all(r["motivo"] == "PERDIDA"
      for r in svc.trazabilidad({"motivo": "PERDIDA"})))
check("filtro por área", all(r["area"] == "Operaciones"
      for r in svc.trazabilidad({"area_id": 1})))
solo_dano = svc.trazabilidad({"motivo": "DANO"})
check("la sustitución referencia la entrega reemplazada",
      len(solo_dano) == 1 and solo_dano[0]["entrega_reemplazada_id"] == botas_ana["entrega_id"])

seccion("6. R2 — EPP vigentes")
vigentes = svc.epp_vigentes({})
ruts = {r["rut"] for r in vigentes}
check("incluye al trabajador sin EPP", "3-3" in ruts)
check("marcado con sin_epp",
      next(r for r in vigentes if r["rut"] == "3-3")["sin_epp"] is True)
check("excluye desvinculados por defecto", "4-4" not in ruts)
check("los incluye con incluir_inactivos",
      "4-4" in {r["rut"] for r in svc.epp_vigentes({"incluir_inactivos": True})})
ana_vig = [r for r in vigentes if r["rut"] == "1-1"]
# El reporte no puede contener ninguna entrega que haya sido reemplazada: sus
# filas con EPP deben coincidir exactamente con las entregas vigentes de los
# trabajadores activos según la BD.
con_epp = [r for r in vigentes if not r["sin_epp"]]
vigentes_bd = escalar(db, """
    SELECT COUNT(*) FROM entregas_epp e
    JOIN personal p ON p.rut = e.rut AND COALESCE(p.activo, TRUE) = TRUE
    WHERE NOT EXISTS (SELECT 1 FROM entregas_epp r WHERE r.entrega_reemplazada_id = e.entrega_id)
""")
check("no lista entregas reemplazadas",
      len(con_epp) == vigentes_bd, f"reporte={len(con_epp)} bd={vigentes_bd}")
check("Ana tiene 2 EPP vigentes (casco repuesto + botas sustituidas)",
      len(ana_vig) == 2, f"{[(r['producto'], r['talla']) for r in ana_vig]}")

filtrado = svc.epp_vigentes({"producto_id": 2})
sin_botas = [r for r in filtrado if r["sin_epp"]]
check("filtrando por producto sigue mostrando a quien NO lo tiene",
      any(r["rut"] == "2-2" for r in sin_botas))

resumen = svc.repo.epp_vigentes_resumen({})
check("resumen: una fila por trabajador",
      len(resumen) == escalar(db, "SELECT COUNT(*) FROM personal WHERE activo"))
check("resumen concatena el detalle",
      any("Casco" in (r["detalle"] or "") for r in resumen))

seccion("7. R3 — Stock y quiebres")
epp.ajustar_stock(3, None, 0, "forzar quiebre", None)      # protector → QUIEBRE
epp.ajustar_stock(1, None, 4, "forzar bajo mínimo", None)  # casco (min 5) → BAJO
stock = svc.stock()
por_prod = {(r["producto"], r["talla"]): r for r in stock}
check("una fila por producto+talla, sin duplicar los sin talla",
      len(stock) == escalar(db, "SELECT COUNT(*) FROM stock_epp"), f"{len(stock)} filas")
check("QUIEBRE detectado", por_prod[("Protector auditivo", None)]["estado"] == "QUIEBRE")
check("BAJO detectado", por_prod[("Casco", None)]["estado"] == "BAJO")
check("OK detectado", por_prod[("Botas de seguridad", "L")]["estado"] == "OK")
check("déficit calculado", por_prod[("Casco", None)]["deficit"] == 1)
check("consumo del casco cuenta las entregas",
      por_prod[("Casco", None)]["consumo"] == 3,
      str(por_prod[("Casco", None)]["consumo"]))
check("cobertura NULL sin consumo",
      por_prod[("Botas de seguridad", "L")]["cobertura_dias"] is None)
check("solo_alertas filtra", all(r["estado"] != "OK" for r in svc.stock(solo_alertas=True)))

seccion("8. Dashboard")
d = svc.dashboard(desde="2026-01-01", hasta="2026-12-31", meses=3)
k = d["kpis"]
check("cuenta unidades, no filas",
      k["entregas_periodo"] == escalar(db, "SELECT SUM(cantidad) FROM entregas_epp"))
check("líneas por separado",
      k["lineas_periodo"] == escalar(db, "SELECT COUNT(*) FROM entregas_epp"))
check("reposiciones por pérdida", k["reposiciones_perdida"] == 1)
check("sustituciones por daño", k["sustituciones_dano"] == 1)
check("productos en quiebre", k["productos_en_quiebre"] == 1)
check("bajo mínimo incluye el quiebre", k["productos_bajo_minimo"] == 2)
check("trabajadores activos", k["trabajadores_activos"] == 3)
check("trabajadores sin EPP", k["trabajadores_sin_epp"] == 1)

serie = d["serie_mensual"]
check("serie con la ventana pedida", len(serie) == 3, f"{[p['mes'] for p in serie]}")
check("meses sin entregas presentes en 0", any(p["total"] == 0 for p in serie))
check("serie ordenada", [p["mes"] for p in serie] == sorted(p["mes"] for p in serie))

d12 = svc.dashboard(desde="2026-01-01", hasta="2026-08-31", meses=12)
check("serie de 12 meses termina en el mes de 'hasta'",
      d12["serie_mensual"][-1]["mes"] == "2026-08")
check("la entrega del borde aparece en julio, no en agosto",
      next(p for p in d12["serie_mensual"] if p["mes"] == "2026-07")["total"] >= 1)

check("por_area agrupa", len(d["por_area"]) >= 1)
check("por_motivo con etiquetas legibles",
      all(p["nombre"] in ("Nueva", "Reposición por pérdida", "Sustitución por daño")
          for p in d["por_motivo"]))
check("alertas_stock solo trae alertas",
      all(r["estado"] != "OK" for r in d["alertas_stock"]))

seccion("9. Excel")
for nombre, buf in (
    ("trazabilidad", svc.trazabilidad_excel({})),
    ("epp vigentes", svc.epp_vigentes_excel({})),
    ("stock", svc.stock_excel()),
):
    data = buf.getvalue()
    check(f"{nombre}: archivo xlsx no vacío", len(data) > 1000 and data[:2] == b"PK",
          f"{len(data)} bytes")

from openpyxl import load_workbook   # noqa: E402
wb = load_workbook(svc.epp_vigentes_excel({}))
check("el Excel de vigentes trae 2 hojas", wb.sheetnames == ["EPP vigentes", "Resumen por trabajador"],
      str(wb.sheetnames))
hoja = wb["EPP vigentes"]
check("encabezado escrito", hoja.cell(1, 1).value == "RUT")
check("filas = filas del reporte", hoja.max_row == len(vigentes) + 1,
      f"{hoja.max_row - 1} vs {len(vigentes)}")

seccion("Resultado")
db.close()
if fallos:
    print(f"\n{len(fallos)} CHECK(S) FALLARON:")
    for f in fallos:
        print(f"  - {f}")
    sys.exit(1)
print("\nTodos los checks pasaron.")
