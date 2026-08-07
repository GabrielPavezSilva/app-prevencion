"""
Motor de generación de Excel — infraestructura sin dominio.

Extraído del ExcelReportBuilder del dominio lavandería (Fase 5): la parte de
estilos/armado era reutilizable, el mapeo de columnas estaba acoplado a prendas
y se descartó. Acá las columnas se declaran por reporte con `Columna`.

Uso:
    (ExcelBuilder()
        .add_sheet("Entregas", filas, [Columna("fecha", "Fecha", ancho=12), ...])
        .build())   -> BytesIO listo para StreamingResponse
"""
import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Any, Callable, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


@dataclass
class Columna:
    """
    Declara una columna del Excel.

    key     — clave del dict de la fila
    label   — encabezado visible
    ancho   — ancho fijo; si es None se calcula del contenido (con tope)
    formato — formato numérico de openpyxl (ej. "dd-mm-yyyy", "#,##0")
    fmt     — transformación del valor antes de escribirlo
    """
    key: str
    label: str
    ancho: Optional[int] = None
    formato: Optional[str] = None
    fmt: Optional[Callable[[Any], Any]] = None


_ANCHO_MAX = 45
_ANCHO_MIN = 8

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="D8DEE4")
_BORDE = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


class ExcelBuilder:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)   # se crean hojas con nombre explícito

    def add_sheet(self, nombre: str, filas: List[Dict[str, Any]],
                  columnas: List[Columna]) -> "ExcelBuilder":
        # Excel no admite nombres de hoja de más de 31 caracteres
        ws = self.wb.create_sheet(nombre[:31])

        ws.append([c.label for c in columnas])
        for celda in ws[1]:
            celda.font = _HEADER_FONT
            celda.fill = _HEADER_FILL
            celda.alignment = _HEADER_ALIGN
            celda.border = _BORDE
        ws.freeze_panes = "A2"

        for fila in filas:
            ws.append([_valor(fila.get(c.key), c) for c in columnas])

        if filas and columnas:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(filas) + 1}"

        for idx, col in enumerate(columnas, start=1):
            letra = get_column_letter(idx)
            if col.formato:
                for celda in ws[letra][1:]:      # se salta el encabezado
                    celda.number_format = col.formato
            ws.column_dimensions[letra].width = col.ancho or _ancho_auto(ws, letra, col.label)

        return self

    def build(self) -> io.BytesIO:
        salida = io.BytesIO()
        self.wb.save(salida)
        salida.seek(0)
        return salida


def _valor(v: Any, col: Columna) -> Any:
    if col.fmt is not None:
        v = col.fmt(v)
    if isinstance(v, bool):
        return "Sí" if v else "No"
    # openpyxl escribe datetime/date nativos; el resto de objetos van como texto
    if v is None or isinstance(v, (int, float, str, datetime, date)):
        return v
    return str(v)


def _ancho_auto(ws, letra: str, label: str) -> int:
    largo = len(label)
    for celda in ws[letra]:
        if celda.value is not None:
            largo = max(largo, len(str(celda.value)))
    return max(_ANCHO_MIN, min(largo + 3, _ANCHO_MAX))
