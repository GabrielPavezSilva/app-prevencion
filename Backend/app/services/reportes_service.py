import io
from datetime import datetime
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from sqlalchemy.orm import Session
from app.repositories.reportes_repository import ReportesRepository
from app.schemas.templates import DEFAULT_TEMPLATES
from app.core.logging_config import logger
import pandas as pd

class ExcelReportBuilder:
    def __init__(self):
        self.wb = Workbook()
        # Eliminar la hoja por defecto si se van a crear nuevas
        self.wb.remove(self.wb.active)
        
        # Estilos comunes
        self.header_font = Font(bold=True, color="FFFFFF")
        self.header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        self.center_alignment = Alignment(horizontal="center", vertical="center")
        self.border = Border(left=Side(style='thin'), 
                             right=Side(style='thin'), 
                             top=Side(style='thin'), 
                             bottom=Side(style='thin'))

    def _style_header(self, ws, row_idx):
        for cell in ws[row_idx]:
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.center_alignment
            cell.border = self.border

    def _auto_adjust_columns(self, ws):
        for column in ws.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2)
            ws.column_dimensions[column[0].column_letter].width = adjusted_width

    def add_dynamic_sheet(self, sheet_name: str, data: List[dict], selected_columns: List[str], column_map: dict):
        # column_map mapea los 'id' del frontend a los nombres de cabecera y keys de diccionario
        # Ej: {'sku': {'label': 'SKU', 'key': 'SKU'}, ...}
        
        ws = self.wb.create_sheet(sheet_name)
        
        # Filtrar solo las columnas que fueron seleccionadas y que existen en este mapeo
        valid_cols = [c for c in selected_columns if c in column_map]
        
        if not valid_cols:
            ws.append(["No hay columnas seleccionadas para esta sección"])
            return self
            
        headers = [column_map[col]['label'] for col in valid_cols]
        ws.append(headers)
        
        for item in data:
            row = [item.get(column_map[col]['key']) for col in valid_cols]
            ws.append(row)
            
        self._style_header(ws, 1)
        self._auto_adjust_columns(ws)
        return self

    def build(self) -> io.BytesIO:
        output = io.BytesIO()
        self.wb.save(output)
        output.seek(0)
        return output

class ReportesService:
    def __init__(self, db: Session):
        self.repository = ReportesRepository(db)
        
        # Mapeo de IDs de frontend a diccionarios de datos de SQLAlchemy
        self.inventario_map = {
            'sku': {'label': 'SKU', 'key': 'SKU'},
            'tipo_prenda': {'label': 'Tipo de Prenda', 'key': 'Tipo'},
            'talla': {'label': 'Talla', 'key': 'Talla'},
            'empresa': {'label': 'Empresa', 'key': 'Empresa'},
            'estado': {'label': 'Estado', 'key': 'Estado'}
        }
        
        self.asignaciones_map = {
            'fecha_entrega': {'label': 'Fecha Entrega', 'key': 'FechaEntrega'},
            'rut_asignado': {'label': 'RUT Asignado', 'key': 'Rut'}
        }

    def generar_reporte_personalizado(self, config: dict):
        builder = ExcelReportBuilder()
        selected_columns = config.get("selectedColumns", [])
        start_date = config.get("dateRange", {}).get("start")
        end_date = config.get("dateRange", {}).get("end")
        
        # Verificar si hay columnas de inventario seleccionadas
        if any(col in self.inventario_map for col in selected_columns):
            inventario_data = self.repository.get_inventario()
            builder.add_dynamic_sheet("Inventario", inventario_data, selected_columns, self.inventario_map)
            
        # Verificar si hay columnas de asignaciones seleccionadas
        if any(col in self.asignaciones_map for col in selected_columns):
            asignaciones_data = self.repository.get_asignaciones_filtradas(start_date, end_date)
            # Para las asignaciones incluiremos internamente datos base importantes si los pidieron
            # Podemos extender asignaciones_map para incluir sku, también presente en esa tabla
            asignaciones_extended_map = {**self.asignaciones_map, 'sku': {'label': 'SKU', 'key': 'SKU'}}
            builder.add_dynamic_sheet("Asignaciones", asignaciones_data, selected_columns, asignaciones_extended_map)
            
        return builder.build()

    def procesar_importacion_excel(self, template_id: str, file_content: bytes, filename: str) -> dict:
        # 1. Obtener y validar el template
        template = next((t for t in DEFAULT_TEMPLATES if t.id == template_id), None)
        if not template:
            raise ValueError(f"Template '{template_id}' no es válido o no existe.")
            
        # 2. Leer el Excel/CSV con pandas
        try:
            if filename.endswith('.csv'):
                df = pd.read_csv(io.BytesIO(file_content))
            else:
                df = pd.read_excel(io.BytesIO(file_content))
        except Exception as e:
            logger.error(f"Error al leer archivo subido: {str(e)}")
            raise ValueError("El archivo está corrupto o no se pudo leer correctamente.")
            
        # 3. Validar columnas requeridas vs DF
        required_cols = [col.name for col in template.columns if col.required]
        df_columns = set(df.columns)
        
        missing_cols = [col for col in required_cols if col not in df_columns]
        if missing_cols:
            raise ValueError(f"El archivo no contiene las columnas requeridas: {', '.join(missing_cols)}")
            
        # 4. Limpiar datos y parsear
        # Eliminar filas vacías
        df = df.dropna(how='all')
        records = df.to_dict(orient='records')
        
        if not records:
            raise ValueError("El archivo no contiene datos para importar.")

        # 5. Mapear claves de DF a fields definidos en el template para pasárselo al Repo
        mapped_records = []
        name_to_field = {col.name: col.field for col in template.columns}
        
        for row in records:
            mapped_row = {}
            for col_name, value in row.items():
                if col_name in name_to_field:
                    # pd.isna verifica NaNs en general de pandas
                    val = None if pd.isna(value) else value
                    mapped_row[name_to_field[col_name]] = val
            mapped_records.append(mapped_row)

        # 6. Delegar bulk inserts al Repository según el template ID
        # Se requiere inyección estructurada por Tipo, Talla, etc.
        if template_id == "inventario_inicial":
            return self.repository.bulk_insert_inventario_inicial(mapped_records)
        elif template_id == "orden_compra":
            # Si tuviésemos logica distinta para Orden de Compra:
            # return self.repository.bulk_insert_orden_compra(mapped_records)
            raise ValueError("La importación de Ordenes de Compra aún no está implementada en backend.")
        else:
            raise ValueError(f"Lógica de importación no definida para {template_id}")

